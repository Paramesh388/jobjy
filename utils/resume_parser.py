import os
import pypdf
import docx
from utils.text_cleaner import clean_resume_text

class ResumeParsingError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass

def extract_text_from_pdf(filepath: str) -> str:
    """
    Extracts text from a PDF file using pypdf.
    Handles multi-page documents and encrypted/damaged files gracefully.
    """
    try:
        reader = pypdf.PdfReader(filepath)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ResumeParsingError("This PDF is password-protected. Please upload an unlocked PDF.")
                
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)
                
        full_text = "\n".join(extracted_pages)
        return clean_resume_text(full_text)
    except ResumeParsingError:
        raise
    except Exception as e:
        raise ResumeParsingError(f"Error reading PDF file: {str(e)}")

def extract_text_from_docx(filepath: str) -> str:
    """
    Extracts text from a DOCX document including paragraphs and tables.
    """
    try:
        doc = docx.Document(filepath)
        text_parts = []
        
        # Extract from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
                
        # Extract from tables (many resumes format skills/education in tables)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
                    
        full_text = "\n".join(text_parts)
        return clean_resume_text(full_text)
    except Exception as e:
        raise ResumeParsingError(f"Error reading DOCX file: {str(e)}")

def extract_text_from_txt(filepath: str) -> str:
    """
    Extracts text from a plain text file, attempting common encodings.
    """
    encodings = ["utf-8", "latin-1", "cp1252", "ascii"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                content = f.read()
                return clean_resume_text(content)
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise ResumeParsingError(f"Error reading TXT file: {str(e)}")
            
    raise ResumeParsingError("Unable to decode the text file using supported character encodings.")

def parse_resume(filepath: str) -> str:
    """
    Dispatcher function to detect file extension and extract text.
    Validates that readable text exists.
    """
    if not os.path.exists(filepath):
        raise ResumeParsingError("The uploaded file does not exist on the server.")
        
    ext = os.path.splitext(filepath)[1].lower()
    
    if ext == ".pdf":
        text = extract_text_from_pdf(filepath)
    elif ext == ".docx":
        text = extract_text_from_docx(filepath)
    elif ext == ".txt":
        text = extract_text_from_txt(filepath)
    else:
        raise ResumeParsingError(f"Unsupported file format: '{ext}'. Please upload a PDF, DOCX, or TXT file.")
        
    if not text or len(text.strip()) < 30:
        raise ResumeParsingError("Unable to extract readable text from this resume. Please upload a text-based PDF/DOCX.")
        
    return text
